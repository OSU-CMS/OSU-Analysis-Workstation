#!/usr/bin/env python3
"""Convert compiled Beamer PDFs into Keynote decks that look exactly like the PDF.

Keynote cannot import a PDF as editable slides, so each PDF page becomes a
full-slide image (Beamer's own rendering: fonts, math, tables, navigation bar).
The slide text is kept as speaker notes so it stays searchable. The slides are
NOT editable text -- for editable slides use make_pptx.py instead.

Pipeline: pdftoppm (page PNGs) -> python-pptx (image-per-slide .pptx, scratch) ->
Keynote via AppleScript (open .pptx, save .key). macOS with Keynote only.

Usage: python beamer_to_keynote.py [--work-dir DIR]     (needs python-pptx, pdftoppm)
"""

import argparse
import shutil
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
DECKS = ["method_overview", "status_update"]
W, H = 13.333, 7.5


def build_pptx(pdf: Path, out: Path, work: Path) -> int:
    prefix = work / pdf.parent.name
    subprocess.run(["pdftoppm", "-r", "200", "-png", str(pdf), str(prefix)], check=True)
    pages = sorted(work.glob(f"{pdf.parent.name}-*.png"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    for number, page in enumerate(pages, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(page), 0, 0, prs.slide_width, prs.slide_height)
        text = subprocess.run(["pdftotext", "-f", str(number), "-l", str(number), "-layout",
                               str(pdf), "-"], capture_output=True, text=True, check=True).stdout
        slide.notes_slide.notes_text_frame.text = text.strip()
    prs.save(out)
    return len(pages)


def to_keynote(pptx: Path, key: Path) -> None:
    script = f'''
with timeout of 900 seconds
  tell application "Keynote"
    set d to open POSIX file "{pptx}"
    delay 4
    save d in POSIX file "{key}"
    close d saving no
  end tell
end timeout'''
    # AppleScript's default 120 s Apple-event timeout is too short for a
    # 20-page image deck, hence the explicit timeout block above.
    subprocess.run(["perl", "-e", "alarm 1000; exec @ARGV", "osascript", "-e", script], check=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--work-dir", type=Path, default=None)
    args = parser.parse_args()
    out_dir = HERE / "keynote"
    out_dir.mkdir(exist_ok=True)
    # Keynote is sandboxed and cannot read the system temp dir, so stage the
    # intermediate .pptx/PNGs under the project and delete them afterwards.
    work = args.work_dir or (out_dir / ".work")
    work.mkdir(parents=True, exist_ok=True)
    try:
        for name in DECKS:
            pdf = HERE / name / "talk.pdf"
            pptx = work / f"{name}_beamer.pptx"
            count = build_pptx(pdf, pptx, work)
            key = out_dir / f"{name}_beamer.key"
            to_keynote(pptx, key)
            print(f"{key.name}: {count} slides")
    finally:
        if args.work_dir is None:
            shutil.rmtree(work, ignore_errors=True)
